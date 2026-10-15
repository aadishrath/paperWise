from io import BytesIO
import asyncio

import pytest
from docx import Document as WordDocument

from app.documents import extract_text, split_text
from app.citation_validation import InvalidCitations, cited_passages
from app.evaluation import score_case, summarize
from app.retrieval import cosine_similarity, rank_chunks
from app.ollama import OllamaError
from app.workflow import AnswerWorkflow


def test_split_text_bounds_chunks_and_preserves_content():
    text = (
        "First sentence. "
        + ("A useful paragraph has enough detail. " * 20)
        + "Last sentence."
    )
    chunks = split_text(text, chunk_size=100, overlap=15)
    assert len(chunks) > 1
    assert all(len(chunk) <= 100 for chunk in chunks)
    assert "First sentence." in chunks[0]
    assert "Last sentence." in chunks[-1]


def test_split_text_rejects_invalid_overlap():
    with pytest.raises(ValueError):
        split_text("Some words", chunk_size=10, overlap=10)


def test_extract_text_supports_plain_text_and_rejects_unknown_types():
    pages = extract_text("notes.txt", b"hello")
    assert len(pages) == 1
    assert pages[0].page_number is None
    assert pages[0].text == "hello"
    with pytest.raises(ValueError, match="Unsupported"):
        extract_text("image.png", b"data")


def test_extract_text_reads_docx_paragraphs_and_tables():
    document = WordDocument()
    document.add_paragraph("Returns are accepted within 30 days.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Support"
    table.cell(0, 1).text = "Weekdays"
    stream = BytesIO()
    document.save(stream)
    pages = extract_text("handbook.docx", stream.getvalue())
    assert "30 days" in pages[0].text
    assert "Support | Weekdays" in pages[0].text


def test_citations_must_reference_retrieved_passages():
    passages = [
        {"filename": "one.txt", "content": "A"},
        {"filename": "two.txt", "content": "B"},
    ]
    selected = cited_passages("The answer uses this source [2].", passages)
    assert [item["citation_number"] for item in selected] == [2]
    with pytest.raises(InvalidCitations, match="not retrieved"):
        cited_passages("This points outside the results [3].", passages)
    with pytest.raises(InvalidCitations, match="did not include"):
        cited_passages("An uncited answer.", passages)


def test_cosine_similarity_and_retrieval_order():
    assert cosine_similarity([1, 0], [1, 0]) == pytest.approx(1.0)
    assert cosine_similarity([1, 0], [0, 1]) == pytest.approx(0.0)
    chunks = [
        {"id": "a", "embedding": [1, 0]},
        {"id": "b", "embedding": [0, 1]},
    ]
    assert [item["id"] for item in rank_chunks([0.9, 0.1], chunks)] == ["a", "b"]


def test_evaluation_scores_source_answer_and_abstention():
    supported = score_case(
        {
            "expected_source": "handbook.pdf",
            "expected_phrases": ["30 days", "delivery"],
        },
        {
            "answer": "Returns are allowed for 30 days after delivery [1].",
            "citations": [{"number": 1, "filename": "handbook.pdf"}],
        },
    )
    abstained = score_case(
        {"expected_source": None, "expected_phrases": []},
        {"answer": "I cannot find this.", "citations": [], "abstained": True},
    )
    assert (
        supported["source_hit"]
        and supported["phrase_recall"] == 1.0
        and supported["citations_valid"]
    )
    assert abstained["abstained"] and abstained["citations_valid"]
    assert summarize([supported, abstained])["cases"] == 2


def test_workflow_reports_missing_question_embedding():
    class EmptyOllama:
        async def embed(self, _texts):
            return []

    with pytest.raises(OllamaError, match="no vector"):
        asyncio.run(
            AnswerWorkflow(database=None, ollama=EmptyOllama()).run("A valid question")
        )


def test_ollama_embedding_payload_rejects_non_finite_values():
    with pytest.raises(OllamaError, match="unexpected embedding"):
        OllamaClient._read_embeddings({"embeddings": [[1.0, float("nan")]]}, expected=1)
